from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import _Cell, _Row, Table
from docx.text.paragraph import Paragraph

from .dates import contract_date_text
from .money import format_money, money_in_words
from .names import buyer_initials, safe_filename


@dataclass(slots=True)
class GeneratedDocument:
    path: Path
    file_name: str
    total_cents: int


class DocumentTemplateError(RuntimeError):
    pass


def _iter_paragraphs(container: DocumentType | _Cell) -> Iterable[Paragraph]:
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _all_paragraphs(doc: DocumentType) -> Iterable[Paragraph]:
    yield from _iter_paragraphs(doc)
    for section in doc.sections:
        for area in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            yield from _iter_paragraphs(area)


def _replace_in_paragraph(paragraph: Paragraph, old: str, new: str) -> int:
    replaced = 0
    while old in "".join(run.text for run in paragraph.runs):
        runs = paragraph.runs
        full = "".join(run.text for run in runs)
        start = full.index(old)
        end = start + len(old)

        positions: list[tuple[int, int]] = []
        cursor = 0
        for index, run in enumerate(runs):
            next_cursor = cursor + len(run.text)
            positions.append((cursor, next_cursor))
            cursor = next_cursor

        start_run = next(i for i, (a, b) in enumerate(positions) if a <= start < b or (a == b == start))
        end_run = next(i for i, (a, b) in enumerate(positions) if a < end <= b)
        start_offset = start - positions[start_run][0]
        end_offset = end - positions[end_run][0]

        prefix = runs[start_run].text[:start_offset]
        suffix = runs[end_run].text[end_offset:]
        runs[start_run].text = prefix + new + suffix
        for index in range(start_run + 1, end_run + 1):
            runs[index].text = ""
        replaced += 1
    return replaced


def _replace_all(doc: DocumentType, old: str, new: str) -> int:
    return sum(_replace_in_paragraph(p, old, new) for p in _all_paragraphs(doc))


def _set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _set_cell_text(cell: _Cell, text: str, *, align: WD_ALIGN_PARAGRAPH | None = None) -> None:
    paragraph = cell.paragraphs[0]
    _set_paragraph_text(paragraph, text)
    if align is not None:
        paragraph.alignment = align
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    # Убираем случайные дополнительные абзацы в ячейке.
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def _find_items_table(doc: DocumentType) -> Table:
    for table in doc.tables:
        header = " | ".join(cell.text for cell in table.rows[0].cells)
        if "Наименование, характеристика" in header and "Гарантия" in header:
            return table
    raise DocumentTemplateError("В шаблоне не найдена таблица товаров")


class ContractDocumentGenerator:
    def __init__(self, template_path: Path, output_dir: Path):
        self.template_path = template_path
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(
        self,
        *,
        contract_number: str,
        contract_date: date,
        buyer_fio: str,
        phone: str,
        address: str,
        items: list[dict],
    ) -> GeneratedDocument:
        if not items:
            raise ValueError("В договоре должен быть хотя бы один товар")

        doc = Document(self.template_path)
        total_cents = sum(
            int(item["quantity"]) * int(item["unit_price_cents"]) for item in items
        )
        total_quantity = sum(int(item["quantity"]) for item in items)
        initials = buyer_initials(buyer_fio)

        replacements = {
            "{{NUM}}": contract_number,
            "{{DATE}}": contract_date_text(contract_date),
            "{{FIO}}": buyer_fio,
            "{{TEL}}": phone,
            "{{ADDR}}": address,
            "{{INIT}}": initials,
            "{{SUM}}": format_money(total_cents),
            "{{WORDS}}": money_in_words(total_cents),
            "{{YEAR}}": str(contract_date.year),
        }
        for marker, value in replacements.items():
            count = _replace_all(doc, marker, value)
            if count == 0:
                raise DocumentTemplateError(f"В шаблоне отсутствует маркер {marker}")

        table = _find_items_table(doc)
        if len(table.rows) < 3:
            raise DocumentTemplateError("Таблица товаров должна содержать заголовок, образец и ИТОГО")

        sample_row = table.rows[1]
        total_row = table.rows[-1]
        sample_xml = deepcopy(sample_row._tr)
        table._tbl.remove(sample_row._tr)

        for index, item in enumerate(items, start=1):
            row_xml = deepcopy(sample_xml)
            total_row._tr.addprevious(row_xml)
            row = _Row(row_xml, table)
            cells = row.cells
            if len(cells) != 5:
                raise DocumentTemplateError("В таблице товаров должно быть 5 столбцов")
            _set_cell_text(cells[0], str(index), align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_text(cells[1], str(item["name"]))
            _set_cell_text(cells[2], str(item["quantity"]), align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_text(
                cells[3],
                format_money(int(item["unit_price_cents"])),
                align=WD_ALIGN_PARAGRAPH.RIGHT,
            )
            _set_cell_text(
                cells[4],
                f"{int(item['warranty_months'])} мес.",
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )

        total_cells = total_row.cells
        _set_cell_text(total_cells[0], "ИТОГО:", align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(total_cells[2], str(total_quantity), align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(total_cells[3], format_money(total_cents), align=WD_ALIGN_PARAGRAPH.RIGHT)
        _set_cell_text(total_cells[4], "", align=WD_ALIGN_PARAGRAPH.CENTER)

        file_name = safe_filename(f"{initials} {contract_number}.docx")
        output_path = self.output_dir / file_name
        doc.core_properties.title = f"Договор {contract_number} — {buyer_fio}"
        doc.save(output_path)
        return GeneratedDocument(path=output_path, file_name=file_name, total_cents=total_cents)
