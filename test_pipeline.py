import sys
import asyncio
from pathlib import Path
from datetime import date
import io

root = Path(__file__).resolve().parent
sys.path.insert(0, str(root))

from api.index import (
    ContractRequestModel,
    ContractItemModel,
    preview_contract,
    generate_contract_file,
    save_and_generate_contract,
    list_contracts,
    get_status,
)
from docx import Document

async def run_tests():
    print("1. Testing get_status...")
    st = await get_status()
    print("Status:", st)
    assert st["status"] == "online"

    print("2. Testing preview_contract...")
    items = [
        ContractItemModel(
            name="\u041a\u043e\u043d\u0434\u0438\u0446\u0438\u043e\u043d\u0435\u0440 Daikin Emura 3 FTXJ20AW/RXJ20A",
            quantity=1,
            unit_price_cents=1370650,
            warranty_months=60
        ),
        ContractItemModel(
            name="Wi-Fi \u043a\u043e\u043d\u0442\u0440\u043e\u043b\u043b\u0435\u0440 BRP069B45",
            quantity=2,
            unit_price_cents=80000,
            warranty_months=12
        )
    ]
    payload = ContractRequestModel(
        buyer_fio="\u041a\u043e\u043d\u043e\u043f\u0430\u0446\u043a\u0438\u0439 \u0418\u043b\u044c\u044f \u0421\u0435\u0440\u0433\u0435\u0435\u0432\u0438\u0447",
        phone="+375 (25) 672-96-37",
        address="\u0433. \u041c\u0438\u043d\u0441\u043a, \u0443\u043b. \u042f\u043d\u043a\u0438 \u041b\u0443\u0447\u0438\u043d\u044b, \u0434. 22, \u043a\u0432. 52",
        contract_date="2026-09-17",
        items=items
    )
    prev = await preview_contract(payload)
    print("Normalized FIO:", prev["buyer_fio"])
    print("Initials:", prev["initials"])
    print("Total words:", prev["total_words"])
    assert prev["buyer_fio"] == "\u041a\u043e\u043d\u043e\u043f\u0430\u0446\u043a\u0438\u0439 \u0418\u043b\u044c\u044f \u0421\u0435\u0440\u0433\u0435\u0435\u0432\u0438\u0447"
    assert prev["initials"] == "\u0418.\u0421.\u041a\u043e\u043d\u043e\u043f\u0430\u0446\u043a\u0438\u0439"
    assert prev["total_cents"] == 1370650 + 2 * 80000

    print("3. Testing generate_contract_file (Word Docx generation)...")
    res = await generate_contract_file(payload)
    doc_bytes = res.body
    print("Generated docx size:", len(doc_bytes), "bytes")
    assert len(doc_bytes) > 10000

    doc = Document(io.BytesIO(doc_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "\u041a\u043e\u043d\u043e\u043f\u0430\u0446\u043a\u0438\u0439 \u0418\u043b\u044c\u044f \u0421\u0435\u0440\u0433\u0435\u0435\u0432\u0438\u0447" in full_text
    assert "{{FIO}}" not in full_text
    assert "{{NUM}}" not in full_text
    print("Word document generated & validated: All placeholders replaced!")

    table = doc.tables[1]
    print("Table rows count:", len(table.rows))
    assert len(table.rows) == 4
    print("Row 1 (Daikin):", [cell.text for cell in table.rows[1].cells])
    print("Row 2 (Wi-Fi):", [cell.text for cell in table.rows[2].cells])

    print("4. Testing save_and_generate_contract & list_contracts...")
    save_result = await save_and_generate_contract(payload)
    print("Save result:", save_result)
    assert save_result["success"] is True

    contracts = await list_contracts(search="\u043a\u043e\u043d\u043e\u043f\u0430\u0446\u043a\u0438\u0439")
    print("Found contracts count:", len(contracts))
    assert len(contracts) > 0
    print("First found contract:", contracts[0]["contract_number"], contracts[0]["buyer_fio"])

    print("\n=== ALL DIRECT TESTS PASSED PERFECTLY! ===")

if __name__ == "__main__":
    asyncio.run(run_tests())
